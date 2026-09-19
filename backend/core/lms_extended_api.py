import io
import json
import mimetypes
import os
import re
import uuid
from datetime import datetime
from pathlib import Path
from urllib.parse import quote, urlparse

import requests
from django.conf import settings
from django.db.models import Avg, Count, Q, Sum
from django.http import FileResponse, HttpResponse, JsonResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.utils.html import strip_tags
from django.views.decorators.http import require_http_methods
from docx import Document

from . import cloud, openedx_bridge
from .lms_models import (
    Course,
    CourseCategory,
    CourseEnrollment,
    CourseEvent,
    CourseRegistrationProfile,
    CourseTag,
    LearningAsset,
    LearningPath,
    Lesson,
    LessonProgress,
    SourceConnection,
)
from .platform_runtime_v3 import core_role, ensure_platform_workspaces
from .pulsar import PulsarError, complete


ZOTERO_ROOT = 'https://api.zotero.org'
SAFE_NAME = re.compile(r'[^A-Za-z0-9._ -]+')
MAX_EVENT_METADATA = 20000


def _json_body(request):
    try:
        data = json.loads(request.body or '{}')
    except (json.JSONDecodeError, UnicodeDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def _error(code, status=400, **extra):
    return JsonResponse({'ok': False, 'error': code, **extra}, status=status)


def _admin(user):
    if not user or not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    spaces = ensure_platform_workspaces(user)
    return core_role(user, spaces['core']) in {'owner', 'admin'}


def _course(course_id):
    return Course.objects.filter(pk=course_id).first()


def _enrollment(user, course):
    if not user or not user.is_authenticated:
        return None
    return CourseEnrollment.objects.filter(
        user=user,
        course=course,
        status__in=[
            CourseEnrollment.Status.ACTIVE,
            CourseEnrollment.Status.PAUSED,
            CourseEnrollment.Status.COMPLETED,
        ],
    ).first()


def _can_access(user, course):
    return _admin(user) or bool(_enrollment(user, course))


def _safe_filename(value):
    name = Path(str(value or '')).name.strip() or 'file'
    return SAFE_NAME.sub('_', name)[:255] or 'file'


def _metadata(value):
    clean = value if isinstance(value, dict) else {}
    encoded = json.dumps(clean, ensure_ascii=False)
    return clean if len(encoded.encode('utf-8')) <= MAX_EVENT_METADATA else {}


def _event(user, course, kind, *, enrollment=None, lesson=None, duration_seconds=0, metadata=None):
    return CourseEvent.objects.create(
        user=user,
        enrollment=enrollment,
        course=course,
        lesson=lesson,
        kind=kind,
        duration_seconds=max(0, min(24 * 60 * 60, int(duration_seconds or 0))),
        metadata=_metadata(metadata),
    )


def _asset_json(item):
    return {
        'id': item.pk,
        'course_id': item.course_id,
        'lesson_id': item.lesson_id,
        'kind': item.kind,
        'title': item.title,
        'original_name': item.original_name,
        'source_url': item.source_url,
        'mime_type': item.mime_type,
        'size': item.size,
        'metadata': item.metadata,
        'download_url': f'/api/lms/assets/{item.pk}/download/' if item.kind == LearningAsset.Kind.FILE else '',
        'created_at': item.created_at.isoformat(),
    }


def _connection_json(item):
    return {
        'id': item.pk,
        'provider': item.provider,
        'label': item.label,
        'library_type': item.library_type,
        'library_id': item.library_id,
        'connected': bool(item.encrypted_token),
        'updated_at': item.updated_at.isoformat(),
    }


def _zotero_url(item):
    root = 'users' if item.library_type == 'user' else 'groups'
    return f'{ZOTERO_ROOT}/{root}/{quote(str(item.library_id), safe="")}'


def _zotero_request(item, endpoint, *, params=None):
    try:
        response = requests.get(
            _zotero_url(item) + endpoint,
            headers={
                'Zotero-API-Key': cloud._decrypt(item.encrypted_token),
                'Zotero-API-Version': '3',
                'Accept': 'application/json',
            },
            params=params or {},
            timeout=(5, 20),
        )
    except requests.RequestException as exc:
        raise ValueError('zotero_unavailable') from exc
    if response.status_code in {401, 403}:
        raise ValueError('zotero_credentials_invalid')
    if response.status_code < 200 or response.status_code >= 300:
        raise ValueError('zotero_request_failed')
    try:
        return response.json()
    except ValueError as exc:
        raise ValueError('zotero_invalid_response') from exc


def _zotero_item_json(row):
    data = row.get('data') if isinstance(row, dict) and isinstance(row.get('data'), dict) else {}
    creators = []
    for creator in data.get('creators') or []:
        if not isinstance(creator, dict):
            continue
        name = creator.get('name') or ' '.join(
            value for value in [creator.get('firstName'), creator.get('lastName')] if value
        )
        if name:
            creators.append(str(name))
    return {
        'key': str(row.get('key') or data.get('key') or ''),
        'title': str(data.get('title') or 'Untitled'),
        'item_type': str(data.get('itemType') or ''),
        'creators': creators[:12],
        'date': str(data.get('date') or ''),
        'publication_title': str(data.get('publicationTitle') or ''),
        'doi': str(data.get('DOI') or ''),
        'url': str(data.get('url') or ''),
        'abstract': str(data.get('abstractNote') or '')[:3000],
    }


@require_http_methods(['POST'])
def lms_event(request, course_id):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    course = _course(course_id)
    if not course:
        return _error('course_not_found', 404)
    enrollment = _enrollment(request.user, course)
    if not enrollment and not _admin(request.user):
        return _error('course_enrollment_required', 403)

    data = _json_body(request)
    kind = str(data.get('kind') or '')
    if kind not in CourseEvent.Kind.values:
        return _error('invalid_event_kind')
    lesson = None
    if data.get('lesson_id'):
        lesson = Lesson.objects.filter(pk=data['lesson_id'], module__course=course).first()
        if not lesson:
            return _error('lesson_not_found', 404)
    event = _event(
        request.user,
        course,
        kind,
        enrollment=enrollment,
        lesson=lesson,
        duration_seconds=data.get('duration_seconds') or 0,
        metadata=data.get('metadata'),
    )
    return JsonResponse({'ok': True, 'event_id': event.pk}, status=201)


@require_http_methods(['GET', 'POST', 'DELETE'])
def lms_registration_profile(request, course_id):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    course = _course(course_id)
    if not course:
        return _error('course_not_found', 404)
    enrollment = _enrollment(request.user, course)
    if not enrollment:
        return _error('course_enrollment_required', 403)
    profile, _ = CourseRegistrationProfile.objects.get_or_create(enrollment=enrollment)

    if request.method == 'GET':
        return JsonResponse({
            'ok': True,
            'schema': course.registration_schema if isinstance(course.registration_schema, list) else [],
            'answers': profile.answers,
            'completed': profile.completed,
        })
    if request.method == 'DELETE':
        profile.delete()
        return JsonResponse({'ok': True})

    data = _json_body(request)
    answers = data.get('answers')
    if not isinstance(answers, dict):
        return _error('answers_required')
    schema = course.registration_schema if isinstance(course.registration_schema, list) else []
    required = [
        str(field.get('key') or '').strip()
        for field in schema
        if isinstance(field, dict) and field.get('required')
    ]
    missing = [key for key in required if key and answers.get(key) in (None, '', [])]
    if missing:
        return _error('required_profile_fields_missing', 409, fields=missing)
    profile.answers = answers
    profile.completed = True
    profile.completed_at = timezone.now()
    profile.save()
    return JsonResponse({'ok': True, 'completed': True})


@require_http_methods(['GET', 'POST', 'DELETE'])
def zotero_connection(request):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    qs = SourceConnection.objects.filter(user=request.user, provider=SourceConnection.Provider.ZOTERO)
    if request.method == 'GET':
        return JsonResponse({'ok': True, 'connections': [_connection_json(item) for item in qs]})
    if request.method == 'DELETE':
        connection_id = request.GET.get('id')
        qs.filter(pk=connection_id).delete()
        return JsonResponse({'ok': True})

    data = _json_body(request)
    library_type = str(data.get('library_type') or 'user').strip().lower()
    if library_type not in {'user', 'group'}:
        return _error('invalid_library_type')
    library_id = str(data.get('library_id') or '').strip()
    api_key = str(data.get('api_key') or '').strip()
    label = str(data.get('label') or 'Zotero').strip()[:120]
    if not library_id or not api_key:
        return _error('library_id_and_api_key_required')

    probe = SourceConnection(
        user=request.user,
        provider=SourceConnection.Provider.ZOTERO,
        label=label or 'Zotero',
        library_type=library_type,
        library_id=library_id,
        encrypted_token=cloud._encrypt(api_key),
    )
    try:
        _zotero_request(probe, '/items', params={'limit': 1, 'format': 'json'})
    except ValueError as exc:
        return _error(str(exc), 400)

    item, _ = SourceConnection.objects.update_or_create(
        user=request.user,
        provider=SourceConnection.Provider.ZOTERO,
        library_type=library_type,
        library_id=library_id,
        defaults={'label': label or 'Zotero', 'encrypted_token': cloud._encrypt(api_key)},
    )
    return JsonResponse({'ok': True, 'connection': _connection_json(item)}, status=201)


@require_http_methods(['GET'])
def zotero_items(request):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    item = SourceConnection.objects.filter(
        pk=request.GET.get('connection_id'),
        user=request.user,
        provider=SourceConnection.Provider.ZOTERO,
    ).first()
    if not item:
        return _error('source_connection_not_found', 404)
    query = str(request.GET.get('q') or '').strip()[:240]
    params = {'limit': min(50, max(1, int(request.GET.get('limit') or 20))), 'format': 'json', 'sort': 'dateModified', 'direction': 'desc'}
    if query:
        params['q'] = query
        params['qmode'] = 'everything'
    try:
        rows = _zotero_request(item, '/items', params=params)
    except ValueError as exc:
        return _error(str(exc), 502)
    return JsonResponse({'ok': True, 'items': [_zotero_item_json(row) for row in rows if isinstance(row, dict)]})


def _source_context(user, connection_id, keys):
    if not connection_id or not keys:
        return []
    item = SourceConnection.objects.filter(pk=connection_id, user=user).first()
    if not item:
        return []
    out = []
    for key in [str(value) for value in keys[:8] if value]:
        try:
            row = _zotero_request(item, f'/items/{quote(key, safe="")}', params={'format': 'json'})
            out.append(_zotero_item_json(row))
        except ValueError:
            continue
    return out


@require_http_methods(['POST'])
def lms_ai_tutor(request, course_id):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    course = _course(course_id)
    if not course:
        return _error('course_not_found', 404)
    enrollment = _enrollment(request.user, course)
    if not enrollment:
        return _error('course_enrollment_required', 403)

    data = _json_body(request)
    question = str(data.get('question') or '').strip()[:6000]
    if not question:
        return _error('question_required')
    lesson = None
    if data.get('lesson_id'):
        lesson = Lesson.objects.filter(pk=data['lesson_id'], module__course=course).first()
        if not lesson:
            return _error('lesson_not_found', 404)

    sources = _source_context(
        request.user,
        data.get('source_connection_id'),
        data.get('source_keys') if isinstance(data.get('source_keys'), list) else [],
    )
    lesson_text = ''
    if lesson:
        lesson_text = strip_tags(lesson.body or '')[:12000]
    source_text = '\n'.join(
        f"- {row['title']} | {', '.join(row['creators'])} | {row['date']} | {row['abstract'][:1000]}"
        for row in sources
    )
    history = []
    for turn in (data.get('history') if isinstance(data.get('history'), list) else [])[-8:]:
        if isinstance(turn, dict):
            role = str(turn.get('role') or '')[:20]
            text = str(turn.get('content') or '').strip()[:1400]
            if role and text:
                history.append(f'{role}: {text}')

    learning_config = course.learning_config if isinstance(course.learning_config, dict) else {}
    guidance_mode = str(learning_config.get('ai_guidance_mode') or 'guided').strip().lower()
    if guidance_mode not in {'hint_only', 'guided', 'full'}:
        guidance_mode = 'guided'
    instructor_prompt = str(learning_config.get('ai_instructor_prompt') or '').strip()[:4000]
    guidance_rule = {
        'hint_only': (
            'Do not provide the final answer, finished derivation, or ready-to-submit solution. '
            'Give the smallest useful hint, ask a diagnostic question, and let the learner do the next step.'
        ),
        'guided': (
            'Prefer hints, Socratic questions and partial scaffolding. '
            'Give a direct answer only after the learner has shown meaningful work or explicitly asks for a final explanation.'
        ),
        'full': (
            'You may provide complete explanations and worked solutions, while still explaining the reasoning and checking understanding.'
        ),
    }[guidance_mode]
    system = (
        'You are the Gravitas+ course tutor. Teach rather than merely answer. '
        'Use the language of the learner. Base factual claims on the supplied course and source context. '
        'When context is insufficient, say so and suggest what to inspect next. '
        + guidance_rule + ' '
        'Never invent a Zotero source or claim the learner completed work they did not complete. '
        + (f'Instructor-specific guidance: {instructor_prompt}' if instructor_prompt else '')
    )
    prompt = (
        f'Course: {course.title}\nCourse summary: {course.summary}\n'
        f'Learner progress: {enrollment.progress_percent}%\n'
        + (f'Lesson: {lesson.title}\nLesson material:\n{lesson_text}\n' if lesson else '')
        + (f'Learner selected sources:\n{source_text}\n' if source_text else '')
        + (f'Recent tutor conversation:\n' + '\n'.join(history) + '\n' if history else '')
        + f'Learner question: {question}'
    )
    try:
        answer = complete(system=system, user=prompt, max_tokens=1400, temperature=0.25)
    except PulsarError as exc:
        return _error(str(exc), 503)

    _event(
        request.user,
        course,
        CourseEvent.Kind.AI_USE,
        enrollment=enrollment,
        lesson=lesson,
        metadata={
            'source_count': len(sources),
            'question_chars': len(question),
            'answer_chars': len(answer),
            'guidance_mode': guidance_mode,
        },
    )
    return JsonResponse({'ok': True, 'answer': answer, 'sources': sources, 'provider': 'cloudflare-workers-ai'})


def _course_export_parts(course):
    yield '#', course.title, course.summary or course.description
    for module in course.modules.prefetch_related('lessons').all():
        yield '##', module.title, module.summary
        for lesson in module.lessons.filter(published=True):
            text = strip_tags(lesson.body or '').strip()
            resource = lesson.content_url or ''
            yield '###', lesson.title, '\n\n'.join(value for value in [lesson.summary, text, resource] if value)


def _latex_escape(value):
    replacements = {
        '\\': r'\textbackslash{}',
        '&': r'\&', '%': r'\%', '$': r'\$', '#': r'\#',
        '_': r'\_', '{': r'\{', '}': r'\}', '~': r'\textasciitilde{}',
        '^': r'\textasciicircum{}',
    }
    return ''.join(replacements.get(ch, ch) for ch in str(value or ''))


@require_http_methods(['GET'])
def lms_course_export(request, course_id, fmt):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    course = _course(course_id)
    if not course:
        return _error('course_not_found', 404)
    enrollment = _enrollment(request.user, course)
    if not enrollment and not _admin(request.user):
        return _error('course_enrollment_required', 403)
    fmt = str(fmt).lower()
    if fmt not in {'md', 'tex', 'docx'}:
        return _error('invalid_export_format')

    parts = list(_course_export_parts(course))
    base_name = re.sub(r'[^A-Za-z0-9_-]+', '-', course.slug or f'course-{course.pk}').strip('-') or f'course-{course.pk}'
    if fmt == 'md':
        body = '\n\n'.join(f'{level} {title}\n\n{text}'.strip() for level, title, text in parts)
        response = HttpResponse(body, content_type='text/markdown; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{base_name}.md"'
    elif fmt == 'tex':
        sections = []
        for level, title, text in parts:
            command = {'#': 'section*', '##': 'section', '###': 'subsection'}.get(level, 'paragraph')
            sections.append(f'\\{command}{{{_latex_escape(title)}}}\n{_latex_escape(text)}')
        body = (
            '\\documentclass{article}\n\\usepackage[utf8]{inputenc}\n'
            '\\usepackage[T1]{fontenc}\n\\begin{document}\n'
            + '\n\n'.join(sections)
            + '\n\\end{document}\n'
        )
        response = HttpResponse(body, content_type='application/x-tex; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{base_name}.tex"'
    else:
        document = Document()
        document.add_heading(course.title, 0)
        if course.summary:
            document.add_paragraph(course.summary)
        for level, title, text in parts[1:]:
            document.add_heading(title, level=1 if level == '##' else 2)
            if text:
                document.add_paragraph(text)
        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        response = FileResponse(
            buffer,
            as_attachment=True,
            filename=f'{base_name}.docx',
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )

    _event(
        request.user,
        course,
        CourseEvent.Kind.EXPORT,
        enrollment=enrollment,
        metadata={'format': fmt},
    )
    return response


@require_http_methods(['GET', 'POST'])
def admin_learning_assets(request, course_id):
    if not _admin(request.user):
        return _error('core_admin_required', 403)
    course = _course(course_id)
    if not course:
        return _error('course_not_found', 404)

    if request.method == 'GET':
        return JsonResponse({'ok': True, 'assets': [_asset_json(item) for item in course.assets.select_related('lesson').all()]})

    title = str(request.POST.get('title') or '').strip()[:240]
    source_url = str(request.POST.get('source_url') or '').strip()[:1800]
    kind = str(request.POST.get('kind') or LearningAsset.Kind.FILE)
    uploaded = request.FILES.get('file')
    lesson = None
    if request.POST.get('lesson_id'):
        lesson = Lesson.objects.filter(pk=request.POST['lesson_id'], module__course=course).first()
        if not lesson:
            return _error('lesson_not_found', 404)
    if kind not in LearningAsset.Kind.values:
        return _error('invalid_asset_kind')
    if kind == LearningAsset.Kind.FILE and not uploaded:
        return _error('file_required')
    if kind != LearningAsset.Kind.FILE and not source_url:
        return _error('source_url_required')
    if source_url:
        parsed = urlparse(source_url)
        if parsed.scheme not in {'http', 'https'} or not parsed.netloc:
            return _error('invalid_source_url')
    if uploaded and (uploaded.size <= 0 or uploaded.size > settings.LMS_ASSET_MAX_BYTES):
        return _error('file_size_invalid', 413, max_bytes=settings.LMS_ASSET_MAX_BYTES)

    item = LearningAsset(
        course=course,
        lesson=lesson,
        kind=kind,
        title=title or (uploaded.name if uploaded else source_url),
        source_url=source_url,
        uploaded_by=request.user,
    )
    if uploaded:
        root = Path(settings.LMS_MEDIA_ROOT) / str(course.pk)
        root.mkdir(parents=True, exist_ok=True)
        name = _safe_filename(uploaded.name)
        path = root / f'{uuid.uuid4().hex}-{name}'
        with path.open('wb') as handle:
            for chunk in uploaded.chunks():
                handle.write(chunk)
        item.original_name = name
        item.storage_path = str(path)
        item.mime_type = (uploaded.content_type or mimetypes.guess_type(name)[0] or '')[:180]
        item.size = uploaded.size
    item.save()
    return JsonResponse({'ok': True, 'asset': _asset_json(item)}, status=201)


@require_http_methods(['GET', 'DELETE'])
def learning_asset_detail(request, asset_id):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    item = LearningAsset.objects.select_related('course').filter(pk=asset_id).first()
    if not item:
        return _error('asset_not_found', 404)
    if request.method == 'DELETE':
        if not _admin(request.user):
            return _error('core_admin_required', 403)
        path = item.storage_path
        item.delete()
        if path:
            try:
                os.remove(path)
            except FileNotFoundError:
                pass
        return JsonResponse({'ok': True})
    if not _can_access(request.user, item.course):
        return _error('course_enrollment_required', 403)
    return JsonResponse({'ok': True, 'asset': _asset_json(item)})


@require_http_methods(['GET'])
def learning_asset_download(request, asset_id):
    if not request.user.is_authenticated:
        return _error('authentication_required', 401)
    item = LearningAsset.objects.select_related('course').filter(pk=asset_id, kind=LearningAsset.Kind.FILE).first()
    if not item or not item.storage_path or not os.path.isfile(item.storage_path):
        return _error('asset_not_found', 404)
    if not _can_access(request.user, item.course):
        return _error('course_enrollment_required', 403)
    return FileResponse(open(item.storage_path, 'rb'), as_attachment=True, filename=item.original_name or item.title)


def _meta_json():
    return {
        'categories': [{
            'id': item.pk,
            'slug': item.slug,
            'name': item.name,
            'description': item.description,
            'position': item.position,
            'active': item.active,
        } for item in CourseCategory.objects.all()],
        'tags': [{'id': item.pk, 'slug': item.slug, 'name': item.name} for item in CourseTag.objects.all()],
    }


@require_http_methods(['GET', 'POST', 'PATCH', 'DELETE'])
def admin_lms_meta(request):
    if not _admin(request.user):
        return _error('core_admin_required', 403)
    if request.method == 'GET':
        return JsonResponse({'ok': True, **_meta_json()})
    data = _json_body(request)
    kind = str(data.get('kind') or '')
    Model = CourseCategory if kind == 'category' else CourseTag if kind == 'tag' else None
    if Model is None:
        return _error('invalid_meta_kind')
    if request.method == 'POST':
        name = str(data.get('name') or '').strip()[:180]
        slug = str(data.get('slug') or '').strip()[:160]
        if not name or not slug:
            return _error('name_and_slug_required')
        defaults = {'name': name}
        if Model is CourseCategory:
            defaults.update(
                description=str(data.get('description') or ''),
                position=max(0, int(data.get('position') or 0)),
                active=bool(data.get('active', True)),
            )
        item, created = Model.objects.update_or_create(slug=slug, defaults=defaults)
        return JsonResponse({'ok': True, 'created': created, **_meta_json()}, status=201 if created else 200)
    try:
        item_id = int(data.get('id'))
    except (TypeError, ValueError):
        return _error('id_required')
    item = Model.objects.filter(pk=item_id).first()
    if not item:
        return _error('not_found', 404)
    if request.method == 'DELETE':
        item.delete()
        return JsonResponse({'ok': True, **_meta_json()})
    if 'name' in data:
        item.name = str(data['name'] or '').strip()[:180]
    if 'slug' in data:
        item.slug = str(data['slug'] or '').strip()[:160]
    if Model is CourseCategory:
        for field in ('description', 'active'):
            if field in data:
                setattr(item, field, data[field] if field == 'active' else str(data[field] or ''))
        if 'position' in data:
            item.position = max(0, int(data['position'] or 0))
    item.save()
    return JsonResponse({'ok': True, **_meta_json()})


def _path_json(item):
    return {
        'id': item.pk,
        'slug': item.slug,
        'title': item.title,
        'summary': item.summary,
        'status': item.status,
        'nodes': item.nodes,
        'edges': item.edges,
        'payment_config': item.payment_config,
        'updated_at': item.updated_at.isoformat(),
    }


@require_http_methods(['GET', 'POST'])
def learning_paths(request):
    admin = _admin(request.user)
    if request.method == 'GET':
        qs = LearningPath.objects.all() if admin and request.GET.get('all') == '1' else LearningPath.objects.filter(status=LearningPath.Status.PUBLISHED)
        return JsonResponse({'ok': True, 'paths': [_path_json(item) for item in qs]})
    if not admin:
        return _error('core_admin_required', 403)
    data = _json_body(request)
    slug = str(data.get('slug') or '').strip()[:190]
    title = str(data.get('title') or '').strip()[:240]
    if not slug or not title:
        return _error('slug_and_title_required')
    if LearningPath.objects.filter(slug=slug).exists():
        return _error('slug_exists', 409)
    status = str(data.get('status') or LearningPath.Status.DRAFT)
    if status not in LearningPath.Status.values:
        return _error('invalid_status')
    nodes = data.get('nodes') if isinstance(data.get('nodes'), list) else []
    edges = data.get('edges') if isinstance(data.get('edges'), list) else []
    item = LearningPath.objects.create(
        slug=slug,
        title=title,
        summary=str(data.get('summary') or ''),
        status=status,
        nodes=nodes,
        edges=edges,
        payment_config=data.get('payment_config') if isinstance(data.get('payment_config'), dict) else {},
        created_by=request.user,
    )
    return JsonResponse({'ok': True, 'path': _path_json(item)}, status=201)


@require_http_methods(['GET', 'PATCH', 'DELETE'])
def learning_path_detail(request, path_id):
    item = LearningPath.objects.filter(pk=path_id).first()
    if not item:
        return _error('path_not_found', 404)
    admin = _admin(request.user)
    if request.method == 'GET':
        if item.status != LearningPath.Status.PUBLISHED and not admin:
            return _error('path_not_found', 404)
        return JsonResponse({'ok': True, 'path': _path_json(item)})
    if not admin:
        return _error('core_admin_required', 403)
    if request.method == 'DELETE':
        item.delete()
        return JsonResponse({'ok': True})
    data = _json_body(request)
    for field in ('slug', 'title', 'summary'):
        if field in data:
            setattr(item, field, str(data[field] or '').strip())
    if 'status' in data:
        if data['status'] not in LearningPath.Status.values:
            return _error('invalid_status')
        item.status = data['status']
    for field in ('nodes', 'edges'):
        if field in data:
            if not isinstance(data[field], list):
                return _error(f'invalid_{field}')
            setattr(item, field, data[field])
    if 'payment_config' in data:
        if not isinstance(data['payment_config'], dict):
            return _error('invalid_payment_config')
        item.payment_config = data['payment_config']
    item.save()
    return JsonResponse({'ok': True, 'path': _path_json(item)})


def _date_filter(qs, request):
    start = str(request.GET.get('from') or '').strip()
    end = str(request.GET.get('to') or '').strip()
    if start:
        parsed = parse_datetime(start)
        if parsed:
            qs = qs.filter(created_at__gte=parsed)
    if end:
        parsed = parse_datetime(end)
        if parsed:
            qs = qs.filter(created_at__lte=parsed)
    return qs


@require_http_methods(['GET'])
def admin_lms_analytics(request):
    if not _admin(request.user):
        return _error('core_admin_required', 403)
    course_id = request.GET.get('course_id')
    user_id = request.GET.get('user_id')
    enrollments = CourseEnrollment.objects.select_related('user', 'course')
    events = CourseEvent.objects.select_related('user', 'course', 'lesson')
    if course_id:
        enrollments = enrollments.filter(course_id=course_id)
        events = events.filter(course_id=course_id)
    if user_id:
        enrollments = enrollments.filter(user_id=user_id)
        events = events.filter(user_id=user_id)
    events = _date_filter(events, request)

    by_kind = {
        row['kind']: {'count': row['count'], 'duration_seconds': row['duration'] or 0}
        for row in events.values('kind').annotate(count=Count('id'), duration=Sum('duration_seconds'))
    }
    lesson_rows = events.filter(lesson__isnull=False).values(
        'lesson_id', 'lesson__title', 'course_id', 'course__title'
    ).annotate(
        views=Count('id', filter=Q(kind=CourseEvent.Kind.LESSON_VIEW)),
        skips=Count('id', filter=Q(kind=CourseEvent.Kind.LESSON_SKIP)),
        dwell_seconds=Sum('duration_seconds', filter=Q(kind=CourseEvent.Kind.LESSON_DWELL)),
        ai_uses=Count('id', filter=Q(kind=CourseEvent.Kind.AI_USE)),
        lab_uses=Count('id', filter=Q(kind=CourseEvent.Kind.LAB_USE)),
    ).order_by('-views', '-dwell_seconds')[:500]

    learner_rows = []
    for enrollment in enrollments[:1000]:
        learner_events = events.filter(user_id=enrollment.user_id, course_id=enrollment.course_id)
        learner_rows.append({
            'enrollment_id': enrollment.pk,
            'user_id': enrollment.user_id,
            'name': enrollment.user.get_full_name() or enrollment.user.email,
            'email': enrollment.user.email,
            'course_id': enrollment.course_id,
            'course': enrollment.course.title,
            'status': enrollment.status,
            'progress_percent': str(enrollment.progress_percent),
            'views': learner_events.filter(kind=CourseEvent.Kind.LESSON_VIEW).count(),
            'skips': learner_events.filter(kind=CourseEvent.Kind.LESSON_SKIP).count(),
            'dwell_seconds': learner_events.filter(kind=CourseEvent.Kind.LESSON_DWELL).aggregate(v=Sum('duration_seconds'))['v'] or 0,
            'ai_uses': learner_events.filter(kind=CourseEvent.Kind.AI_USE).count(),
            'lab_uses': learner_events.filter(kind=CourseEvent.Kind.LAB_USE).count(),
        })

    return JsonResponse({
        'ok': True,
        'summary': {
            'enrollments': enrollments.count(),
            'active': enrollments.filter(status=CourseEnrollment.Status.ACTIVE).count(),
            'completed': enrollments.filter(status=CourseEnrollment.Status.COMPLETED).count(),
            'average_progress': str(enrollments.aggregate(v=Avg('progress_percent'))['v'] or 0),
            'events': events.count(),
            'by_kind': by_kind,
        },
        'learners': learner_rows,
        'lessons': list(lesson_rows),
    })


@require_http_methods(['GET', 'POST'])
def admin_openedx_status(request):
    if not _admin(request.user):
        return _error('core_admin_required', 403)
    if request.method == 'GET':
        return JsonResponse({
            'ok': True,
            'openedx': openedx_bridge.health(),
            'lms_url': settings.OPENEDX_LMS_URL,
            'cms_url': settings.OPENEDX_CMS_URL,
        })
    data = _json_body(request)
    course = _course(data.get('course_id'))
    if not course:
        return _error('course_not_found', 404)
    if course.provider != Course.Provider.OPENEDX or not course.openedx_course_key:
        return _error('openedx_course_mapping_required', 409)
    try:
        details = openedx_bridge.course_details(course.openedx_course_key)
    except openedx_bridge.OpenEdXError as exc:
        return _error(str(exc), 502)
    return JsonResponse({'ok': True, 'course_details': details})
